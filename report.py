"""
PME Proposals Toolkit
----------------------
Consolida dados de propostas comerciais PME (planos de saude PME/PF) a partir
de um CSV e gera um relatorio Excel formatado, com resumo por operadora,
resumo por status e um grafico, agilizando o acompanhamento das propostas.

Uso:
    python report.py propostas_exemplo.csv relatorio.xlsx
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="2F6690", end_color="2F6690", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1C4E6F")
CURRENCY_FMT = 'R$ #,##0.00'


def load_data(csv_path: str) -> pd.DataFrame:
    df = pd.read_csv(csv_path)
    required = {"cliente", "operadora", "vidas", "valor_estimado", "status", "data_envio"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Colunas faltando no CSV: {sorted(missing)}")
    return df


def style_header(ws, row, n_cols):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def autofit(ws, df, start_col=1):
    for i, col in enumerate(df.columns, start=start_col):
        width = max(len(str(col)), df[col].astype(str).map(len).max() if len(df) else 0) + 3
        ws.column_dimensions[get_column_letter(i)].width = width


def write_dataframe(ws, df, start_row=1, start_col=1, currency_cols=None):
    currency_cols = currency_cols or []
    for j, col in enumerate(df.columns, start=start_col):
        ws.cell(row=start_row, column=j, value=col)
    style_header(ws, start_row, len(df.columns))

    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, col in enumerate(df.columns, start=start_col):
            cell = ws.cell(row=i, column=j, value=row[col])
            if col in currency_cols:
                cell.number_format = CURRENCY_FMT
    return start_row + len(df) + 1


def build_report(df: pd.DataFrame, output_path: str):
    wb = Workbook()

    # --- Sheet 1: Propostas (dados brutos) ---
    ws_raw = wb.active
    ws_raw.title = "Propostas"
    write_dataframe(ws_raw, df, currency_cols={"valor_estimado"})
    autofit(ws_raw, df)

    # --- Sheet 2: Resumo por operadora ---
    resumo_op = (
        df.groupby("operadora")
        .agg(qtd_propostas=("id", "count"), total_vidas=("vidas", "sum"), valor_total=("valor_estimado", "sum"))
        .reset_index()
        .sort_values("valor_total", ascending=False)
    )

    ws_summary = wb.create_sheet("Resumo por Operadora")
    ws_summary["A1"] = "Resumo de propostas PME por operadora"
    ws_summary["A1"].font = TITLE_FONT
    next_row = write_dataframe(ws_summary, resumo_op, start_row=3, currency_cols={"valor_total"})
    autofit(ws_summary, resumo_op)

    chart = BarChart()
    chart.title = "Valor total estimado por operadora"
    chart.y_axis.title = "R$"
    chart.x_axis.title = "Operadora"
    data = Reference(ws_summary, min_col=4, min_row=3, max_row=2 + len(resumo_op))
    cats = Reference(ws_summary, min_col=1, min_row=4, max_row=2 + len(resumo_op))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws_summary.add_chart(chart, f"A{next_row + 2}")

    # --- Sheet 3: Resumo por status ---
    resumo_status = (
        df.groupby("status")
        .agg(qtd_propostas=("id", "count"), valor_total=("valor_estimado", "sum"))
        .reset_index()
        .sort_values("qtd_propostas", ascending=False)
    )
    ws_status = wb.create_sheet("Resumo por Status")
    ws_status["A1"] = "Resumo de propostas por status"
    ws_status["A1"].font = TITLE_FONT
    write_dataframe(ws_status, resumo_status, start_row=3, currency_cols={"valor_total"})
    autofit(ws_status, resumo_status)

    wb.save(output_path)


def main():
    if len(sys.argv) < 3:
        print("Uso: python report.py <entrada.csv> <saida.xlsx>")
        sys.exit(1)

    csv_path, output_path = sys.argv[1], sys.argv[2]
    if not Path(csv_path).exists():
        print(f"Arquivo nao encontrado: {csv_path}")
        sys.exit(1)

    df = load_data(csv_path)
    build_report(df, output_path)
    print(f"Relatorio gerado em: {output_path}")


if __name__ == "__main__":
    main()
